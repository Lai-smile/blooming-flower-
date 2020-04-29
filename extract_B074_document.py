import os
import zipfile
import re
from threading import Thread
import numpy as np
import pandas as pd
import docx
from docx import Document  # 导入库
from win32com import client as wc
from win32com.client import Dispatch, constants
from docx import Document
from threading import Thread

base_url = "F:\B074"
unzip_url = "F:\B074_unzip"  # 一次解压后的新文件保存的路径
unzip_high_url = "F:\B074_originalI_file"  # 二次解压后的新文件保存的路径
dirs = os.listdir(base_url)
print(dirs)
print(len(dirs))

def un_zip(zip, dst_dir, basic_path):  # zip:单个压缩包名，dst_dir：目标文件夹，basic_path：压缩包的目录部分
    """unzip zip file"""
    fz = zipfile.ZipFile(os.path.join(basic_path, zip), 'r')
    for file in fz.namelist():
        fz.extract(file, dst_dir)

def un_zip_to_doc(zip, dst_dir, basic_path):  # zip:单个压缩包名，dst_dir：目标文件夹，basic_path：压缩包的目录部分
    """unzip zip file"""
    zip_list = []
    fz = zipfile.ZipFile(os.path.join(basic_path, zip), 'r')
    for file in fz.namelist():
        fz.extract(file, dst_dir)
        if file.split('.')[1] == 'doc':
            zip_dict={zip:file}
        else:
            zip_dict={zip:None}
        zip_list.append(zip_dict)
    return zip_list


# def parse_doc(f):
#     """读取doc
#     """
#     w = wc.Dispatch('Word.Application')
#     doc = w.Documents.Open(FileName=f)
#     t = doc.Tables[0]  # 根据文件中的图表选择信息
#     name = t.Rows[0].Cells[1].Range.Text
#     situation = t.Rows[0].Cells[5].Range.Text
#     people = t.Rows[1].Cells[1].Range.Text
#     title = t.Rows[1].Cells[3].Range.Text
#     print(name, situation, people, title)
#     doc.Close()


customer_code_list = []
for filename in dirs:
    customer_code = filename.split('.')[0][2:5]
    customer_code_list.append(customer_code)


one_zip_list = os.listdir(unzip_url)
two_zip_list = []
doc_list = []
content_list = []
result_dict = {}
content_dict = {}

# for zip in dirs:
#     one_zip_list = un_zip(zip, unzip_url, base_url)  # 一次解压后得到的原始文件名列表
for original_zip in dirs:
    for two_zip in one_zip_list:
        two_zip_last_name=two_zip.split('.')[0].split(' ')[-1]
        if two_zip_last_name==original_zip.split('.')[0]:
            two_zip_list.append(two_zip) # '原始文件名'列表

for third_zip in two_zip_list:  #third_zip（原始文件，zip格式）
    doc_list=un_zip_to_doc(third_zip,unzip_high_url,unzip_url)  #二次解压得到‘原始文件名’压缩包和‘要提取的doc文件”的字典式列表
print(doc_list)
etract_list=[]  #要提取的文件名表
requirements_list=[]   #技术要求与工程要求内容表
for doc in doc_list:
    etract_list.append(doc.values())
    if os.path.isfile(unzip_high_url + str(doc.values())):
        w = wc.Dispatch('Word.Application')
        document = w.Documents(unzip_high_url + str(doc))  # 读入文件
        result_dict['doc'] = document
        tables = document.tables  # 获取文件中的表格集
        table = tables[0]  # 获取文件中的第一个表格
        WorksRequirements_0 = table.cell(5, 0).text  # 工程要求或技术要求,cell(5,0)表示第6行第1列数据，以此类推
        WorksRequirements_1 = table.cell(5, 1).text  # 工程要求或技术要求内容
        specialThings_0 = table.cell(6, 0).text  # 注意事项
        specialThings_1 = table.cell(6, 1).text  # 注意事项内容
        content = {WorksRequirements_0: WorksRequirements_1, specialThings_0: specialThings_1}
        content_dict = {doc:content}
    else:
        content_dict = {doc:None}
    content_list.append(content_dict)

for content in content_list:
    requirements_list.append(content.values())
# threads = []
# t1 = Thread(target=zip_exract,args=(one_zip_list,))
# threads.append(t1)
# t2 = Thread(target=zip_exract,args=(one_zip_list,))
# threads.append(t2)

# if __name__=="__main__":
#     for t in threads:
#         t.setDaemon(True)
#         t.start()
#     t.join()
data = {
    '客户代码': customer_code_list,
    '原始订单名': dirs,
    '原始文件名': one_zip_list,
    '要提取的文件名': etract_list,
    '提取项': requirements_list,

}
data_df = pd.DataFrame(
    np.arange(6315).reshape(1263, 5),
    # columns=[
    #     '客户代码',
    #     '原始订单名',
    #     '原始文件名',
    #     '要提取的文件名',
    #     '提取项',
    # ]
)
print(one_zip_list)
print(etract_list)
print(requirements_list)
data_df.insert(0,'客户代码',customer_code_list)
data_df.insert(1,'原始订单名',dirs)
data_df.insert(2,'原始文件名',one_zip_list)
data_df.insert(3,'要提取的文件名',etract_list)
data_df.insert(4,'提取项',requirements_list)
print(data_df)
writer = pd.ExcelWriter('BO74_data.xlsx')
data_df.to_excel(writer, float_format='%.5f')
writer.save()
