import os
from threading import Thread
import numpy as np
import pandas as pd
import docx
from docx import Document  # 导入库
from win32com import client as wc
from win32com.client import Dispatch, constants
from docx import Document
from decompress_fold import decompress_folder

base_url = "F:\B074"
unzip_url = r"F:\B074_unzip\\"  # 一次解压后的新文件保存的路径
unzip_high_url = "F:\B074_originalI_file"  # 二次解压后的新文件保存的路径
dirs = os.listdir(base_url) #原始订单名列表
srcond_dirs = os.listdir(unzip_url)

customer_code_list = [] #客户代码
for filename in dirs:
    customer_code = filename.split('.')[0][2:5]
    customer_code_list.append(customer_code)

one_zip_path = [] #原始文件名列表
original_order_list = []   #要提取的文件名列表
for one_zip in srcond_dirs:
    for dir in dirs:
        if one_zip.split(' ')[-1] ==dir:
             one_zip_path.append(decompress_folder(one_zip))

full_path=''
for file_path in one_zip_path:
    full_path=os.path.join(unzip_url,file_path)
    original_order_list.append(os.listdir(file_path)[-1])
content_list=[] #提取项
for doc in original_order_list:
    with open(os.path.join(full_path,doc)) as f:
        w = wc.Dispatch('Word.Application')
        document = w.Documents(unzip_high_url + str(doc))  # 读入文件
        tables = document.tables  # 获取文件中的表格集
        table = tables[0]  # 获取文件中的第一个表格
        WorksRequirements_0 = table.cell(5, 0).text  # 工程要求或技术要求,cell(5,0)表示第6行第1列数据，以此类推
        WorksRequirements_1 = table.cell(5, 1).text  # 工程要求或技术要求内容
        specialThings_0 = table.cell(6, 0).text  # 注意事项
        specialThings_1 = table.cell(6, 1).text  # 注意事项内容
        content = {WorksRequirements_0: WorksRequirements_1, specialThings_0: specialThings_1}
        content_dict = {doc: content}
        content_list.append(content_dict)

data_df = pd.DataFrame(
    np.arange(6315).reshape(1263, 5),
    # columns=[
    #     '客户代码',
    #     '原始订单名',
    #     '原始文件名',
    #     '要提取的文件名',
    #     '提取项',
    # ]
)
print(one_zip_path)
print(original_order_list)
print(content_list)
data_df.insert(0,'客户代码',customer_code_list)
data_df.insert(1,'原始订单名',dirs)
data_df.insert(2,'原始文件名',one_zip_path)
data_df.insert(3,'要提取的文件名',original_order_list)
data_df.insert(4,'提取项',content_list)
print(data_df)
writer = pd.ExcelWriter('BO74_data.xlsx')
data_df.to_excel(writer, float_format='%.5f')
writer.save()